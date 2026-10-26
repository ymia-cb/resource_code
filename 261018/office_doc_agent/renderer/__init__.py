"""Stage 4b：IR → OOXML（Word）。

约束：
- 只用 python-docx 逐元素写入，禁止手拼 <w:p>/<c:f> 字符串。
- 渲染器抛「结构非法」异常，而非静默产出损坏文件。
- 显式设置中文字体（宋体）——正式文档必须显式声明字体，否则导出 PDF 出方框。
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from ..interfaces import IR, Renderer

# 默认中文字体：宋体（font-family 第一项必须是中文字体，避免 PDF 方框）
_CJK_FONT = "宋体"


class DocxRenderer:
    """把中间表示渲染为 Word 文档。"""

    def render(self, ir: IR, out_path: Path) -> Path:
        doc = Document()

        for block in ir.get("blocks", []):
            btype = block.get("type")
            if btype == "paragraph":
                self._add_paragraph(doc, block)
            elif btype == "table":
                self._add_table(doc, block)
            elif btype == "formula_cell":
                # Word 端公式单元格：MVP 阶段降级为文本提示（Excel 公式由 recalc 层处理）
                self._add_formula_placeholder(doc, block)
            else:
                raise ValueError(f"未知 block 类型: {btype}")

        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out_path))
        return out_path

    @staticmethod
    def _set_cjk_font(run) -> None:
        """给 run 显式设置中文字体（w:eastAsia），避免 PDF 方框。"""
        run.font.name = _CJK_FONT
        # 同时设置 eastAsia 字体（Word 中文实际用的是这个）
        run._element.rPr.rFonts.set(qn("w:eastAsia"), _CJK_FONT)

    @staticmethod
    def _add_paragraph(doc, block: dict) -> None:
        text = block.get("text", "")
        # 占位符绑定替换：{key} → value
        for key, value in block.get("bindings", {}).items():
            text = text.replace("{" + key + "}", str(value))
        para = doc.add_paragraph()
        run = para.add_run(text)
        DocxRenderer._set_cjk_font(run)

    @staticmethod
    def _add_table(doc, block: dict) -> None:
        rows = block.get("rows", [])
        if not rows:
            return
        n_cols = max(len(r.get("cells", [])) for r in rows)
        table = doc.add_table(rows=len(rows), cols=n_cols)
        table.style = block.get("style", "Table Grid")
        for i, row in enumerate(rows):
            for j in range(n_cols):
                cell_data = row.get("cells", [])[j] if j < len(row.get("cells", [])) else {}
                text = cell_data.get("text", "")
                for key, value in cell_data.get("bindings", {}).items():
                    text = text.replace("{" + key + "}", str(value))
                cell = table.cell(i, j)
                cell.text = ""  # 清空默认段落
                run = cell.paragraphs[0].add_run(text)
                DocxRenderer._set_cjk_font(run)

    @staticmethod
    def _add_formula_placeholder(doc, block: dict) -> None:
        formula = block.get("formula", "")
        p = doc.add_paragraph()
        run = p.add_run(f"[公式] {formula}（缓存值由重算层回填）")
        run.font.size = Pt(9)
        DocxRenderer._set_cjk_font(run)
